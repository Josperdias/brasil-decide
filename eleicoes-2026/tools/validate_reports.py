#!/usr/bin/env python3
"""Valida os relatórios gerados pelo app no emulador (dist/report.*): cada formato precisa abrir em leitores independentes."""
import csv, io, json, sys, zipfile
d = sys.argv[1] if len(sys.argv) > 1 else "dist"
bad = []

def check(name, fn):
    try:
        info = fn()
        print(f"OK   {name}: {info}")
    except Exception as e:
        print(f"FALHA {name}: {type(e).__name__}: {e}")
        bad.append(name)

def pdf():
    from pypdf import PdfReader
    r = PdfReader(f"{d}/report.pdf")
    txt = "\n".join((p.extract_text() or "") for p in r.pages)
    assert len(r.pages) >= 1 and "Brasil Decide" in txt, "sem texto esperado"
    return f"{len(r.pages)} páginas, {len(txt)} caracteres"

def docx():
    import docx
    doc = docx.Document(f"{d}/report.docx")
    assert any("Brasil Decide" in p.text for p in doc.paragraphs), "sem título"
    return f"{len(doc.paragraphs)} parágrafos, {len(doc.tables)} tabelas"

def xlsx():
    import openpyxl
    wb = openpyxl.load_workbook(f"{d}/report.xlsx")
    rows = {ws.title: ws.max_row for ws in wb.worksheets}
    assert "Resumo" in rows
    return str(rows)

def csv_():
    raw = open(f"{d}/report.csv", "rb").read().decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(raw), delimiter=";"))
    assert len(rows) > 5, "poucas linhas"
    return f"{len(rows)} linhas"

def json_():
    o = json.load(open(f"{d}/report.json", encoding="utf-8"))
    assert o["titulo"] and isinstance(o["tabelas"], list)
    return f"{len(o['tabelas'])} tabelas, {len(o['destaques'])} destaques"

def png():
    from PIL import Image
    im = Image.open(f"{d}/report.png")
    im.verify()
    return f"{im.size}"

for n, f in (("pdf", pdf), ("docx", docx), ("xlsx", xlsx), ("csv", csv_), ("json", json_), ("png", png)):
    check(n, f)
sys.exit(1 if bad else 0)
